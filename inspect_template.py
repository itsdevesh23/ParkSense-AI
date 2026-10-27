import docx

doc = docx.Document('Journal_Manuscript_Template with Formatting Instructions.docx')
with open('template_dump.txt', 'w', encoding='utf-8') as f:
    f.write(f'PARAGRAPHS ({len(doc.paragraphs)}):\n')
    for i, p in enumerate(doc.paragraphs):
        f.write(f'[{i}] ({p.style.name}): {p.text}\n')
    f.write(f'\nTABLES ({len(doc.tables)}):\n')
    for t_idx, table in enumerate(doc.tables):
        f.write(f'\n--- TABLE {t_idx} ({len(table.rows)}x{len(table.columns)}) ---\n')
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.replace('\n', ' ') for c in row.cells]
            f.write(f'Row {r_idx}: {" | ".join(cells)}\n')

print('Dumped successfully!')
