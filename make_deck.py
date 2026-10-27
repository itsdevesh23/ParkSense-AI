import sys, os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_refined_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Academic & Professional Palette
    DARK_BG = RGBColor(15, 23, 42)        # Slate 900
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)    # Slate 50
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(203, 213, 225) # Slate 300
    PRIMARY = RGBColor(30, 58, 138)       # Blue 900
    SECONDARY = RGBColor(2, 132, 199)     # Sky 600
    ACCENT_TEAL = RGBColor(13, 148, 136)  # Teal 600
    TEXT_MAIN = RGBColor(30, 41, 59)      # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    LIGHT_BLUE = RGBColor(239, 246, 255)  # Blue 50
    LIGHT_TEAL = RGBColor(240, 253, 250)  # Teal 50

    TOTAL_SLIDES = 11

    def add_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, rubric_tag, title_text, subtitle_text):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.733), Inches(1.15))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = rubric_tag.upper()
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = SECONDARY
        p0.font.name = 'Calibri'
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(21)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY
        p1.font.name = 'Calibri'
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(10.5)
            p2.font.color.rgb = TEXT_MUTED
            p2.font.name = 'Calibri'

    def add_footer(slide, current_page):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(11.733), Inches(0.38))
        tf = footer_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f'Review 1: Problem Formulation & Design (40 Marks) | Team 4 | Federated Multi-Lot Smart Parking | Slide {current_page} of {TOTAL_SLIDES}'
        p.font.size = Pt(9)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = 'Calibri'

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    # ==================== SLIDE 1: TITLE SLIDE ====================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG
    bg1.line.fill.background()

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.15), Inches(11.333), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = 'PROJECT REVIEW 1 • 40 MARKS (PROBLEM FORMULATION & DESIGN)'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)

    p = tf1.add_paragraph()
    p.text = 'Federated Multi-Lot Smart Parking Forecasting\nUsing Attention-Based LSTM on IoT Streams'
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(12)
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = 'Collaborative Short-Term Availability Forecasting with Decentralized Data Locality and Client Personalization'
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(203, 213, 225)

    # Team details card with exact names and register numbers
    details_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.65), Inches(11.333), Inches(2.4))
    dtf = details_box.text_frame
    dtf.word_wrap = True

    p = dtf.paragraphs[0]
    p.text = 'TEAM DETAILS: TEAM NUMBER 4'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)

    p = dtf.add_paragraph()
    p.text = 'Member 1: Devesh Munagala (23BCT0149)   |   Member 2: Abhinav (23BCT0142)'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(4)

    p = dtf.add_paragraph()
    p.text = 'Domain: Smart Cities • IoT Sensor Networks • Federated Deep Learning • Time-Series Forecasting'
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_before = Pt(6)

    p = dtf.add_paragraph()
    p.text = 'Target Horizon: 30–60 Minute Ahead Availability Prediction  |  Primary Testbed: Birmingham Parking (UCI)'
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(148, 163, 184)

    p = dtf.add_paragraph()
    p.text = 'Sustainable Development Goals: SDG 11 (Sustainable Cities) & SDG 9 (Infrastructure & Innovation)'
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(52, 211, 153)

    # ==================== SLIDE 2: CONTEXT, IOT & SDGS ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s2)
    add_header(s2, 'Rubric Criterion 1 • 4 Marks', 'Problem Context, IoT Environment & SDG Alignment', 'Real-world urban parking dilemmas, decentralized IoT streaming, and justified sustainable impact.')
    add_footer(s2, 2)

    col_w = Inches(3.64)
    col_gap = Inches(0.4)
    top_pos = Inches(1.65)
    card_h = Inches(5.0)

    create_card(s2, Inches(0.8), top_pos, col_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '1. Urban Congestion & IoT Setup'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    
    bullets = [
        'Cruising Congestion: Drivers searching for vacant parking bays contribute substantially to downtown traffic congestion, travel delays, and excess fuel consumption (Shoup, 2006).',
        'Continuous IoT Streams: In-ground magnetometers, barrier gates, and loop detectors emit continuous timestamped occupancy readings.',
        'Short-Horizon Demand: Static occupancy signs become stale quickly; 30–60 minute predictive horizons are required for effective driver guidance and demand management.'
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = '• ' + b
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(8)

    create_card(s2, Inches(0.8) + col_w + col_gap, top_pos, col_w, card_h)
    tb = s2.shapes.add_textbox(Inches(1.0) + col_w + col_gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '2. The Data-Sharing Dilemma'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    
    bullets = [
        'Multi-Operator Fragmentation: Urban parking facilities are operated by independent municipal, commercial, and transit authorities.',
        'Data Governance Constraints: Commercial confidentiality, competitive boundaries, and data-governance policies frequently limit raw sensor data pooling between operators.',
        'Centralized vs. Federated: Centralized training requires pooling raw historical streams onto one server; federated learning offers an alternative where training occurs locally.'
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = '• ' + b
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(8)

    create_card(s2, Inches(0.8) + (col_w + col_gap)*2, top_pos, col_w, card_h, bg_color=LIGHT_TEAL, border_color=ACCENT_TEAL)
    tb = s2.shapes.add_textbox(Inches(1.0) + (col_w + col_gap)*2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '3. Justified SDG Alignment'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    
    bullets = [
        'SDG 11: Sustainable Cities & Communities\nTarget 11.2 & 11.6: Accurate short-term availability forecasting may help reduce cruising mileage, urban congestion, and associated vehicular emissions by providing timely availability signals.',
        'SDG 9: Industry, Innovation & Infrastructure\nTarget 9.4 & 9.5: Applies decentralized machine learning across existing municipal IoT infrastructure, supporting collaborative predictive models without requiring raw data centralization.'
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = '• ' + b
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(8)

    # ==================== SLIDE 3: LITERATURE SURVEY & SYNTHESIS ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s3)
    add_header(s3, 'Rubric Criterion 2 • 10 Marks (Part 1)', 'Literature Survey & Thematic Synthesis (30 Papers)', 'Systematic review of 30 peer-reviewed works meeting the 70% recency floor across 5 core research themes.')
    add_footer(s3, 3)

    card_w2 = Inches(5.66)
    gap2 = Inches(0.4)

    create_card(s3, Inches(0.8), top_pos, card_w2, card_h)
    tb = s3.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), card_w2 - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '5 Core Thematic Silos & Recency Audit'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    themes = [
        'Survey Corpus Audit: Full literature matrix comprises 30 peer-reviewed works (detailed in Review 1 Report). 22 published in 2021–2026 (73.3% recency, exceeding the 70% rubric floor); 8 identified as seminal foundational works (exempted under rubric).',
        'Theme A: Smart Parking Forecasting (Stolfi 2017, Zhao 2022, Xiao 2023, Zhang MAT-LSTM 2024) — Transition from shallow fits to attention sequence deep learning; predominantly centralized.',
        'Theme B: Attention Time-Series DL (Vaswani 2017, Li 2022, Kemik 2024, Paliari 2021) — Temporal attention dynamically weights crucial transition intervals over standard recurrent units.',
        'Theme C: Federated Learning for IoT (McMahan 2017, Li FedProx 2020, Pandya 2023, Alterkawi & Dib 2025) — Decentralized model aggregation; identifies statistical heterogeneity and lack of benchmarks as open gaps.',
        'Theme D: Federated Time-Series & Mobility (Huang FedParking 2021, Sharma & Kaur 2024, Attiah FedPark 2025, Rocco et al. 2025) — Collaborative FL proven viable; FedParking pioneered vehicular edge FL for parking estimation.',
        'Theme E: Intelligent Transportation Systems (Chen & Zhang 2024) — Situates privacy-preserving decentralized learning within urban mobility systems.'
    ]
    for t in themes:
        p = tf.add_paragraph()
        p.text = '• ' + t
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    create_card(s3, Inches(0.8) + card_w2 + gap2, top_pos, card_w2, card_h, bg_color=LIGHT_BLUE, border_color=PRIMARY)
    tb = s3.shapes.add_textbox(Inches(1.0) + card_w2 + gap2, top_pos + Inches(0.2), card_w2 - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Synthesis: Solved vs. Open Questions'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    synthesis = [
        'EVIDENCE ESTABLISHED IN PRIOR WORK:\n1. Deep sequence networks (LSTM/GRU) capture non-linear diurnal patterns better than static statistical baselines.\n2. Attention mechanisms improve short-term forecasting by focusing on informative historical sequence segments.\n3. FedAvg can collaboratively aggregate models across distributed IoT nodes without transferring raw data.',
        'LIMITATIONS IN THE SURVEYED LITERATURE:\n1. Parking studies focus predominantly on centralized datasets or single facilities.\n2. Federated mobility models primarily target highway networks (e.g., PeMS road graphs) rather than multi-lot parking streams.\n3. Lot-to-lot demand differences (e.g., commercial vs. transit lots) can degrade a single global model; personalization in smart parking remains under-investigated.'
    ]
    for s in synthesis:
        p = tf.add_paragraph()
        p.text = s
        p.font.size = Pt(10.2)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(8)

    # ==================== SLIDE 4: LITERATURE COMPARISON TABLE ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s4)
    add_header(s4, 'Rubric Criterion 2 • 10 Marks (Part 2)', 'Literature Comparison Table (Representative Anchor Papers)', 'Summary comparison matrix. Full 30-paper comparison matrix is documented in the project report / appendix.')
    add_footer(s4, 4)

    table_shape = s4.shapes.add_table(7, 6, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.0))
    table = table_shape.table

    table.columns[0].width = Inches(1.4)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(1.5)
    table.columns[3].width = Inches(1.6)
    table.columns[4].width = Inches(2.4)
    table.columns[5].width = Inches(2.433)

    headers = ['Author & Year', 'Paper Title', 'Method', 'Dataset / Domain', 'Key Contribution / Finding', 'Technical Limitation']
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.05)
        cell.margin_bottom = Inches(0.05)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    table_data = [
        ('Stolfi et al. (2017)', 'Predicting Car Park Occupancy Rates in Smart Cities', 'Polynomial / Fourier', 'Birmingham Parking (UCI)', 'Demonstrated short-term predictability of municipal parking occupancy streams.', 'Linear/shallow models struggle with non-linear demand shifts; single-operator focus.'),
        ('Huang et al. (2021)', 'FedParking: FL-Based Parking Space Estimation with Parked Vehicles', 'Federated LSTM', 'Parked Vehicle Testbed', 'Pioneered federated learning for parking space estimation via edge nodes (IEEE TVT).', 'Lacks temporal attention; no client personalization; omitted 6-tier comparative benchmark.'),
        ('Zhang et al. (2024)', 'Parking Space Availability Prediction Using MAT-LSTM', 'Multi-Head Attention LSTM', 'Smart parking sensor streams', 'Multi-head attention yields 48% accuracy gain over standard LSTM on parking (MDPI IJGI).', 'Requires centralized raw data pooling; lacks federated multi-operator aggregation.'),
        ('Sharma & Kaur (2024)', 'Fog-Based Federated Time Series Forecasting for IoT Data', 'Fog Federated Forecaster', 'Decentralized IoT logs', 'Evaluated federated learning on non-IID IoT edge sensor streams (Springer JNSM).', 'Generic IoT benchmark; does not investigate parking domain personalization.'),
        ('Rocco et al. (2025)', 'Fedflow: Personalized Federated Learning for Passenger Flow', 'Personalized FL (Fedflow)', 'Multi-operator transit', 'Local fine-tuning helped bridge operational differences across operators (Springer ML).', 'Applied to passenger flow; untested on continuous parking bay sensor streams.'),
        ('Alterkawi & Dib (2025)', 'Federated Learning for Smart Cities: A Thematic Review', 'Thematic Review (116 studies)', 'Smart-city FL literature', 'Identified client heterogeneity and absence of standardized benchmarks as #1 open gaps (MDPI FI).', 'Survey article; does not implement an empirical parking model.')
    ]

    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BLUE if (i % 2 == 1) else WHITE
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(8.8)
            p.font.color.rgb = TEXT_MAIN
            if j == 1:
                p.font.italic = True

    # ==================== SLIDE 5: RESEARCH GAPS ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s5)
    add_header(s5, 'Rubric Criterion 3 • 6 Marks', 'Traceable Research Gaps & Proposed Technical Responses', 'Specific technical gaps derived from the surveyed literature matrix and mapped to our proposed responses.')
    add_footer(s5, 5)

    gaps = [
        ('Gap 1: Single-Facility Focus in Parking Literature', 'Traced to: Stolfi (2017), Dahiya (2021), Zhao MePark (2022) [P5, P11, P12]', 'The surveyed parking forecasting literature primarily evaluates models on isolated facilities or single centralized systems.', 'Proposed Response: Simulate a collaborative multi-lot framework across distributed parking-lot clients partitioned by SystemCodeNumber.'),
        ('Gap 2: Reliance on Centralized Data Pooling', 'Traced to: Huang (2022), Zeng (2022), Zhang MAT-LSTM (2024) [P13, P14, P23]', 'Prior machine learning approaches typically require pooling raw occupancy records centrally.', 'Proposed Response: Apply FedAvg to aggregate model updates while keeping raw sensor records strictly local.'),
        ('Gap 3: Limited Attention Integration in Federated Parking', 'Traced to: Huang FedParking (2021), Li (2022), Attiah FedPark (2025) [P9, P16, P27]', 'FedParking (P9) demonstrated vehicular edge FL for parking but omitted temporal attention; MAT-LSTM (P23) proved attention boosts accuracy by 48% but remains centralized.', 'Proposed Response: Design an Attention-LSTM client forecaster to weight informative historical time steps under federated aggregation.'),
        ('Gap 4: Inter-Lot Demand Heterogeneity (Non-IID Patterns)', 'Traced to: Fallah Per-FedAvg (2020), Sharma & Kaur (2024), Rocco Fedflow (2025) [P8, P25, P29]', 'Different parking facilities exhibit distinct demand distributions, causing client drift that degrades a single global model.', 'Proposed Response: Incorporate local fine-tuning to personalize the aggregated model for individual parking lots.'),
        ('Gap 5: Need for Comprehensive Multi-Model Benchmarking', 'Traced to: Paliari (2021), Xiao (2023), Alterkawi & Dib (2025) [P10, P20, P30]', 'Existing studies lack unified comparisons across statistical, ML, centralized DL, and federated models under identical splits.', 'Proposed Response: Benchmark ARIMA, XGBoost, LSTM, Centralized Attention-LSTM, Federated, and Personalized FL under standardized metrics.')
    ]

    gap_card_h = Inches(0.9)
    gap_card_gap = Inches(0.12)
    start_top = Inches(1.65)

    for idx, (title, trace, desc, resp) in enumerate(gaps):
        top_y = start_top + idx * (gap_card_h + gap_card_gap)
        create_card(s5, Inches(0.8), top_y, Inches(11.733), gap_card_h, bg_color=WHITE, border_color=CARD_BORDER)
        tb = s5.shapes.add_textbox(Inches(1.0), top_y + Inches(0.08), Inches(11.333), gap_card_h - Inches(0.16))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f'{title}   [{trace}]'
        p.font.size = Pt(10.8)
        p.font.bold = True
        p.font.color.rgb = PRIMARY
        
        p2 = tf.add_paragraph()
        p2.text = f'Unsolved Issue: {desc}  -->  {resp}'
        p2.font.size = Pt(9.6)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_before = Pt(2)

    # ==================== SLIDE 6: PROBLEM STATEMENT, OBJECTIVES & SCOPE ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s6)
    add_header(s6, 'Rubric Criterion 4 • 5 Marks', 'Problem Statement, Measurable Objectives & Scope', 'Precise one-sentence problem formulation, 4 achievable semester goals, and strict scope boundary declarations.')
    add_footer(s6, 6)

    create_card(s6, Inches(0.8), Inches(1.65), Inches(11.733), Inches(1.3), bg_color=LIGHT_BLUE, border_color=SECONDARY)
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(1.75), Inches(11.333), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'FORMAL PROBLEM STATEMENT (EXACT ONE-SENTENCE DEFINITION)'
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p2 = tf.add_paragraph()
    p2.text = '“Existing parking availability forecasting approaches often rely on centralized data pooling or isolated single-facility models, creating commercial and governance data-sharing barriers; therefore, a privacy-preserving federated attention-based LSTM framework with local personalization is required to forecast 30–60 minute parking availability across distributed IoT streams while keeping raw occupancy data local.”'
    p2.font.size = Pt(11.5)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY
    p2.space_before = Pt(4)

    create_card(s6, Inches(0.8), Inches(3.1), Inches(6.8), Inches(3.6))
    tb = s6.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(6.4), Inches(3.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '4 Measurable Semester Objectives'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    objs = [
        'Objective 1 (Model Design): Develop an Attention-LSTM neural network to capture temporal dependencies for 30–60 minute ahead parking availability forecasting.',
        'Objective 2 (Federated Architecture): Implement a decentralized FedAvg framework across simulated parking lots, aggregating model parameters without transferring raw occupancy data.',
        'Objective 3 (Personalization): Implement local fine-tuning on client edge nodes to evaluate whether personalization mitigates inter-lot demand heterogeneity.',
        'Objective 4 (Benchmarking): Compare all 6 models: (1) ARIMA, (2) XGBoost, (3) LSTM, (4) Centralized Attention-LSTM, (5) Federated Attention-LSTM, and (6) Personalized Federated Attention-LSTM across MAE, RMSE, MAPE, communication rounds, and the federated–centralized error gap.'
    ]
    for o in objs:
        p = tf.add_paragraph()
        p.text = '• ' + o
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    create_card(s6, Inches(7.8), Inches(3.1), Inches(4.733), Inches(3.6), bg_color=LIGHT_TEAL, border_color=ACCENT_TEAL)
    tb = s6.shapes.add_textbox(Inches(8.0), Inches(3.2), Inches(4.333), Inches(3.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Declared Scope Boundaries'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    scopes = [
        'STRICTLY IN SCOPE:\n• 30–60 minute multi-horizon forecasting.\n• Multi-client federated simulation using real IoT logs.\n• Attention-LSTM local client models.\n• Comparative evaluation across the 6 declared models.\n• Communication round convergence analysis.',
        'EXPLICITLY OUT OF SCOPE (NON-GOALS):\n• Physical IoT sensor hardware deployment.\n• Commercial payment gateway integration.\n• Bay reservation & license-plate recognition.\n• Turn-by-turn vehicle navigation systems.\n• Cryptographic differential privacy / SMPC.'
    ]
    for s in scopes:
        p = tf.add_paragraph()
        p.text = s
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(6)

    # ==================== SLIDE 7: SYSTEM ARCHITECTURE ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s7)
    add_header(s7, 'Rubric Criterion 5 • 8 Marks (Part 1)', 'Proposed System Architecture & Block Diagram', 'Decentralized multi-tier framework linking IoT data feeds, edge Attention-LSTM training, and central FedAvg aggregation.')
    add_footer(s7, 7)

    arch_w = Inches(11.733)
    create_card(s7, Inches(0.8), Inches(1.65), arch_w, Inches(1.0), bg_color=WHITE, border_color=SECONDARY)
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(1.75), arch_w - Inches(0.4), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'TIER 1: DISTRIBUTED IoT SENSOR STREAMS (INDEPENDENT CLIENTS)'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = SECONDARY
    p2 = tf.add_paragraph()
    p2.text = 'Parking Lot 1 (Commuter Lot)  |  Parking Lot 2 (Retail Center)  |  Parking Lot 3 (Municipal)  ...  Parking Lot K\n• Continuous sensor streams: Occupancy counts, capacity, timestamps  • Raw sensor data remains strictly local at each client'
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MAIN

    create_card(s7, Inches(0.8), Inches(2.8), arch_w, Inches(1.3), bg_color=LIGHT_BLUE, border_color=PRIMARY)
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(2.9), arch_w - Inches(0.4), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'TIER 2: CLIENT EDGE PROCESSING & LOCAL ATTENTION-LSTM TRAINING'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    p2 = tf.add_paragraph()
    p2.text = '• Preprocessing: Chronological sliding-window sequences [t - T_in : t]  • Normalization to Availability Ratio in [0, 1]\n• Neural Model: LSTM sequence layers generate hidden state representations  • Temporal attention computes alignment weights alpha_t\n• Local Optimization: Client trains for E local epochs on private loss to update local parameters w_k^(t+1)'
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MAIN

    create_card(s7, Inches(0.8), Inches(4.25), arch_w, Inches(1.2), bg_color=WHITE, border_color=ACCENT_TEAL)
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(4.35), arch_w - Inches(0.4), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'TIER 3: FEDERATED COORDINATION SERVER (WEIGHT AGGREGATION - FedAvg)'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    p2 = tf.add_paragraph()
    p2.text = '• Uplink Transmission: Participating clients send model weights w_k only (No raw occupancy records transmitted)\n• Aggregation Engine: Server computes sample-weighted average:  w_(t+1) = Sum (n_k / n) * w_k^(t+1)\n• Downlink Broadcast: Updated global parameter vector w_(t+1) broadcasted back to all participating clients'
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MAIN

    create_card(s7, Inches(0.8), Inches(5.6), arch_w, Inches(1.1), bg_color=LIGHT_TEAL, border_color=ACCENT_TEAL)
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(5.7), arch_w - Inches(0.4), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'TIER 4: LOCAL CLIENT PERSONALIZATION & SHORT-TERM AVAILABILITY FORECAST'
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL
    p2 = tf.add_paragraph()
    p2.text = '• Personalization: Local fine-tuning adapts the received global model to lot-specific demand patterns\n• Output Horizon: Generates 30–60 minute ahead occupancy and space availability forecasts with MAE/RMSE benchmarking'
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MAIN

    # ==================== SLIDE 8: COMPARATIVE METHODOLOGY & FEDAVG ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s8)
    add_header(s8, 'Rubric Criterion 5 • 8 Marks (Part 2)', 'Comparative Methodology & Algorithmic Workflow', 'Rigorous 6-model benchmarking protocol alongside mathematical steps of Federated Averaging and personalization.')
    add_footer(s8, 8)

    create_card(s8, Inches(0.8), top_pos, card_w2, card_h)
    tb = s8.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), card_w2 - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = '6-Tier Comparative Model Protocol'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    models = [
        'Model 1: Classical Statistical Baseline — ARIMA(p, d, q)\nCaptures linear autoregressive trends; provides traditional reference point.',
        'Model 2: Machine Learning Baseline — XGBoost\nGradient-boosted decision trees trained on engineered temporal/lag features.',
        'Model 3: Standard Deep Sequence Baseline — Plain LSTM\nRecurrent neural network capturing non-linear sequential dependencies.',
        'Model 4: Centralized DL Reference — Centralized Attention-LSTM\nAll data pooled centrally; provides a reference benchmark for accuracy.',
        'Model 5: Proposed Decentralized Model — Federated Attention-LSTM\nFedAvg aggregation over distributed clients; keeps raw sensor logs local.',
        'Model 6: Proposed Personalized Model — Personalized Fed-Attention-LSTM\nLocal fine-tuning on client edge nodes to adapt to lot heterogeneity.'
    ]
    for m in models:
        p = tf.add_paragraph()
        p.text = '• ' + m
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    create_card(s8, Inches(0.8) + card_w2 + gap2, top_pos, card_w2, card_h, bg_color=WHITE, border_color=SECONDARY)
    tb = s8.shapes.add_textbox(Inches(1.0) + card_w2 + gap2, top_pos + Inches(0.2), card_w2 - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Algorithmic Workflow of FedAvg & Personalization'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = SECONDARY

    algo_steps = [
        'Step 1 (Global Init): Server initializes global parameter vector w_0 and broadcasts to all K simulated parking clients.',
        'Step 2 (Local Edge Update): In communication round t, client k trains for E local epochs of SGD on private loss L_k(w; D_k):\n   w_k^(t+1) <-- w_t - eta * Grad(L_k)',
        'Step 3 (Uplink Transmission): Each client sends its updated weights w_k^(t+1) to server. Zero raw records are shared.',
        'Step 4 (Federated Aggregation): Server computes weighted average:\n   w_(t+1) = Sum_{k=1}^K (n_k / n) * w_k^(t+1)',
        'Step 5 (Broadcast & Convergence): Server transmits w_(t+1) back. Repeats for T rounds until validation error converges.',
        'Step 6 (Personalization Phase): Client fine-tunes the global model on its private training data for few local epochs.'
    ]
    for a in algo_steps:
        p = tf.add_paragraph()
        p.text = '• ' + a
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    # ==================== SLIDE 9: EXPERIMENTAL TESTBED & DATASETS ====================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s9)
    add_header(s9, 'Rubric Criterion 5 / 6', 'Experimental Testbed, Datasets & Preprocessing', 'Directly verified UCI sensor parameters, planned preprocessing steps, and target formulation.')
    add_footer(s9, 9)

    create_card(s9, Inches(0.8), top_pos, col_w, card_h)
    tb = s9.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Primary Testbed: Birmingham (UCI)'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    birmingham_pts = [
        'Source: UCI ML Repository (Stolfi et al., 2017).',
        'Direct CSV Verification: Exactly 30 unique parking facilities (SystemCodeNumber) verified from dataset.csv; 35,717 records; 4 raw attributes.',
        'Observed Characteristics: Daytime operating hours (07:46 to 16:30 daily) from Oct 4 to Dec 19, 2016. Nominal 30-min sampling with empirical variations (27–33 min intervals).',
        'Planned Preprocessing Steps:\n1. Regular Grid Resampling: Resampling observations onto an exact uniform 30-min grid.\n2. Anomaly Cleansing: Handling sensor artifacts (e.g. negative occupancy readings).\n3. Imputation: Forward-fill/linear imputation for missing daytime interval gaps.',
        'Client Partition: Maps each of the 30 car parks to an independent federated client.'
    ]
    for b in birmingham_pts:
        p = tf.add_paragraph()
        p.text = '• ' + b
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    create_card(s9, Inches(0.8) + col_w + col_gap, top_pos, col_w, card_h)
    tb = s9.shapes.add_textbox(Inches(1.0) + col_w + col_gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Prediction Target & Secondary Data'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    sec_pts = [
        'Formal Prediction Target:\nAvailability Ratio: y_t = (Capacity - Occupancy) / Capacity in [0, 1].\nModels train directly on continuous ratio y_t to normalize varying lot capacities (ranging 220 to 4,675 spaces).',
        'Physical Output Conversion:\nPer-lot vacant spaces: Free = y_t * Capacity (reported for practical operational interpretability).',
        'Secondary Reference Datasets:\n• Melbourne On-Street: Bay magnetometer events.\n• SFpark: Sensor logs with dynamic occupancy.\nUsed as reference testbeds to validate sensor aggregation pipelines and examine cross-city generalizability.'
    ]
    for s in sec_pts:
        p = tf.add_paragraph()
        p.text = '• ' + s
        p.font.size = Pt(9.6)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(6)

    create_card(s9, Inches(0.8) + (col_w + col_gap)*2, top_pos, col_w, card_h, bg_color=LIGHT_BLUE, border_color=SECONDARY)
    tb = s9.shapes.add_textbox(Inches(1.0) + (col_w + col_gap)*2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Preprocessing & Data Partitioning'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = SECONDARY

    prep_pts = [
        'Sliding Window Generation: Input window of past T_in steps (e.g. 12 steps = 6 hrs) mapped to output horizon T_out (1–2 steps = 30–60 min ahead).',
        'Chronological Splitting:\n• Train Set: 70% (Chronological past)\n• Validation Set: 15%\n• Test Set: 15% (Unseen future)\nStrict zero-lookahead split prevents future data leakage across all clients.',
        'Edge Client Mapping: Each lot acts as an isolated client training locally on its chronological partition.'
    ]
    for pr in prep_pts:
        p = tf.add_paragraph()
        p.text = '• ' + pr
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(8)

    # ==================== SLIDE 10: METRICS, HYPOTHESES & DELIVERABLES ====================
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s10)
    add_header(s10, 'Rubric Criterion 6 • 4 Marks', 'Evaluation Metrics, Research Hypotheses & Deliverables', 'Quantitative evaluation formulas, zero-handling, qualitative testable hypotheses, and deliverables.')
    add_footer(s10, 10)

    create_card(s10, Inches(0.8), top_pos, col_w, card_h)
    tb = s10.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Evaluation Metrics & Formulations'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    metric_pts = [
        'Primary Loss (Ratio in [0, 1]): MAE and RMSE evaluated directly on y_t across all time steps, independent of lot scale.',
        'Per-Lot Space Count Metric: Because capacities differ, physical error is calculated per lot k:\n   MAE_spaces^(k) = MAE_ratio^(k) * Capacity_k\nOverall space error is reported as the mean across lots: (1/K) Sum MAE_spaces^(k).',
        'Zero-Handled MAPE (%): To prevent division-by-zero when a lot is full (y_t = 0), MAPE is smoothed using epsilon = 0.05:\n   MAPE = (100%/N) Sum (|y_t - y_hat_t| / max(y_t, epsilon)).',
        'Variance & Comm. Rounds: Std. deviation of error across lots; rounds to converge.',
        'Error Gap: Delta_FC = MAE(Fed) - MAE(Cent) reported in absolute and relative terms.'
    ]
    for m in metric_pts:
        p = tf.add_paragraph()
        p.text = '• ' + m
        p.font.size = Pt(9.2)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    create_card(s10, Inches(0.8) + col_w + col_gap, top_pos, col_w, card_h)
    tb = s10.shapes.add_textbox(Inches(1.0) + col_w + col_gap, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Testable Research Hypotheses'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    hypo_pts = [
        'Hypothesis 1 (Attention Mechanism):\nAttention-LSTM may reduce forecasting error compared with plain LSTM by dynamically weighting informative historical time steps during peak transition periods.',
        'Hypothesis 2 (Federated Feasibility):\nFederated Attention-LSTM may approach the forecasting accuracy of the centralized reference model while maintaining raw data locality.',
        'Hypothesis 3 (Personalization Efficacy):\nLocal client fine-tuning may improve forecasting accuracy and reduce error disparity for heterogeneous parking facilities compared to pure global FedAvg.'
    ]
    for h in hypo_pts:
        p = tf.add_paragraph()
        p.text = '• ' + h
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(6)

    create_card(s10, Inches(0.8) + (col_w + col_gap)*2, top_pos, col_w, card_h, bg_color=LIGHT_TEAL, border_color=ACCENT_TEAL)
    tb = s10.shapes.add_textbox(Inches(1.0) + (col_w + col_gap)*2, top_pos + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Deliverables & Tools'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    deliv_pts = [
        'Deliverable 1: Modular PyTorch simulation codebase implementing the 6 comparative models.',
        'Deliverable 2: Comprehensive empirical benchmark report documenting MAE, RMSE, MAPE, rounds, and per-lot performance.',
        'Deliverable 3: Research paper manuscript prepared in IEEE conference template targeting Review 2 publication marks.',
        'Tools & Feasibility:\nPython 3.11, PyTorch, Scikit-learn, Statsmodels (ARIMA), Matplotlib/Seaborn. Fully feasible on available workstation hardware.'
    ]
    for d in deliv_pts:
        p = tf.add_paragraph()
        p.text = '• ' + d
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(6)

    # ==================== SLIDE 11: REFERENCES / BIBLIOGRAPHY ====================
    s11 = prs.slides.add_slide(blank_layout)
    add_slide_bg(s11)
    add_header(s11, 'Review 1 Documentation', 'Selected Key References (12 Representative Anchors)', 'Verified publication metadata. Complete 30-paper bibliography provided in Project Report & Appendix.')
    add_footer(s11, 11)

    ref_w = Inches(5.66)
    ref_gap = Inches(0.4)

    # Col 1: Parking & Time-Series
    create_card(s11, Inches(0.8), top_pos, ref_w, card_h)
    tb = s11.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), ref_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Smart Parking & Sequence Forecasting'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    refs_col1 = [
        '[1] D. H. Stolfi, E. Alba, and X. Yao, "Predicting Car Park Occupancy Rates in Smart Cities," in Smart-CT, Springer LNCS, 2017, pp. 107–117.',
        '[2] A. Vaswani et al., "Attention is all you need," in Proc. Advances in Neural Information Processing Systems (NeurIPS), 2017, pp. 5998–6008.',
        '[3] I. Paliari, A. Karanikola, and S. Kotsiantis, "A comparison of optimized LSTM, XGBOOST and ARIMA in Time Series," in Proc. IEEE IISA, 2021, pp. 1–7.',
        '[4] Z. Li, H. Xu, X. Gao et al., "Fusion Attention Mechanism Bidirectional LSTM for Traffic Flow," J. Intell. Transp. Syst., vol. 26, pp. 481–496, 2022.',
        '[5] F. Zhang, K. Shang, L. Yan et al., "Prediction of Parking Space Availability Using Improved MAT-LSTM," ISPRS Int. J. Geo-Inf., 13(5), art. 151, 2024.',
        '[6] H. Kemik, T. Dalyan, and M. Aydogan, "Smart Solutions for Mega-Cities: Utilizing LSTM and MHA in Parking," ISPRS Int. J. Geo-Inf., 13(12), art. 449, 2024.'
    ]
    for r in refs_col1:
        p = tf.add_paragraph()
        p.text = r
        p.font.size = Pt(8.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    # Col 2: Federated Learning & Distributed IoT
    create_card(s11, Inches(0.8) + ref_w + ref_gap, top_pos, ref_w, card_h)
    tb = s11.shapes.add_textbox(Inches(1.0) + ref_w + ref_gap, top_pos + Inches(0.2), ref_w - Inches(0.4), card_h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Federated Learning & Distributed IoT'
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    refs_col2 = [
        '[7] B. McMahan et al., "Communication-efficient learning of deep networks from decentralized data," in Proc. AISTATS, PMLR, 2017, pp. 1273–1282.',
        '[8] A. Fallah, A. Mokhtari, and A. Ozdaglar, "Personalized Federated Learning (Per-FedAvg)," in Proc. NeurIPS, 2020, pp. 3557–3568.',
        '[9] X. Huang, P. Li, R. Yu et al., "FedParking: A Federated Learning Based Parking Space Estimation," IEEE Trans. Veh. Technol., 70(9), pp. 9355–9368, 2021.',
        '[10] M. Sharma and P. Kaur, "Fog-based Federated Time Series Forecasting for IoT Data," J. Netw. Syst. Manage. (Springer), vol. 32, art. 26, 2024.',
        '[11] F. Rocco di Torrepadula et al., "Fedflow: A Personalized FL Framework for Passenger Flow," Machine Learning (Springer), vol. 114, art. 166, 2025.',
        '[12] L. Alterkawi and F. K. Dib, "Federated Learning for Smart Cities: A Thematic Review," Future Internet (MDPI), vol. 17, no. 12, art. 545, 2025.'
    ]
    for r in refs_col2:
        p = tf.add_paragraph()
        p.text = r
        p.font.size = Pt(8.8)
        p.font.color.rgb = TEXT_MAIN
        p.space_before = Pt(5)

    output_path = 'Review_1_Presentation.pptx'
    try:
        prs.save(output_path)
        print(f'Successfully rebuilt refined {output_path} with {TOTAL_SLIDES} slides')
    except PermissionError:
        fallback_path = 'Review_1_Presentation_Updated.pptx'
        prs.save(fallback_path)
        print(f'Notice: {output_path} is currently open in PowerPoint. Saved updated deck as {fallback_path}!')

if __name__ == '__main__':
    build_refined_presentation()
