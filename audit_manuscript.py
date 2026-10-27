"""
audit_manuscript.py
Audits the compiled Journal Manuscript docx file.
"""

import docx
import re
import sys

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document("Journal_Manuscript_Final_Submission.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

total_words = 0
abstract_words = 0
in_abstract = False
headings = []

placeholders = []
placeholder_patterns = [r"\[INSERT", r"\[TODO", r"\[TBD", r"\[X\]", r"SCOPUS-INDEXED JOURNAL SPECIFICATION"]

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    words = len(text.split())
    total_words += words
    
    if text == "Abstract":
        in_abstract = True
        continue
    if in_abstract and "Keywords:" in text:
        in_abstract = False
    elif in_abstract:
        abstract_words += words

    # Detect true section headings (e.g. "1. Introduction", "2. Literature Review", etc.)
    if re.match(r"^([1-6]\.\s+[A-Z]|[1-6]\.[1-4]\s+[A-Z]|References|Declarations)", text):
        headings.append(text)
        
    for pat in placeholder_patterns:
        if re.search(pat, text, re.IGNORECASE):
            placeholders.append((i, text))

print(f"Total Document Word Count: {total_words}")
print(f"Abstract Word Count: {abstract_words}")
print("\nMajor Section Headings Detected:")
for h in headings:
    print(f"  • {h}")

print(f"\nTables Audit ({len(doc.tables)} tables):")
for idx, tbl in enumerate(doc.tables):
    hdr = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    print(f"  Table {idx+1}: {len(tbl.rows)} rows x {len(tbl.columns)} cols | Header: {' | '.join(hdr[:4])}...")

print(f"\nPlaceholders or suspicious flags found: {len(placeholders)}")
for idx, ptext in placeholders:
    print(f"  P[{idx}]: {ptext[:100]}")

refs = [p.text for p in doc.paragraphs if re.match(r"^\[\d+\]", p.text.strip())]
print(f"\nVerified References Count: {len(refs)}")
for r in refs[:3]:
    print(f"  {r[:90]}...")
print("  ...")
for r in refs[-3:]:
    print(f"  {r[:90]}...")
