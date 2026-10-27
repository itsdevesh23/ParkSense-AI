import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

def apply_table_formatting(tbl):
    if len(tbl.rows) > 0:
        header_tr = tbl.rows[0]._tr.get_or_add_trPr()
        header_tr.append(parse_xml(r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

def create_prep_guide():
    src_doc = docx.Document('Review_1_Report.docx.bak')
    out_doc = docx.Document()

    # Page setup
    for sec in out_doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Document Title
    p_title = out_doc.add_paragraph()
    r_title = p_title.add_run("PROJECT REVIEW 1: VIVA PREPARATION & AUDIT GUIDE")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = out_doc.add_paragraph()
    r_sub = p_sub.add_run("Federated Multi-Lot Smart Parking Forecasting Using Attention-Based LSTM on IoT Streams\n")
    r_sub.font.size = Pt(13)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(30, 58, 138)
    r_sub2 = p_sub.add_run("Team 4: Devesh Munagala (23BCT0149) & Abhinav (23BCT0142) | Comprehensive Defence & Viva Handbook\n")
    r_sub2.font.size = Pt(10.5)
    r_sub2.font.color.rgb = RGBColor(100, 116, 139)

    out_doc.add_paragraph("Note: This document contains internal preparation checklists, presentation support slide notes, comprehensive viva Q&A with answers, rubric audits, and the Review 2 execution sequence for student reference during review defense.\n")

    # 1. Rubric & Readiness Checklist (Table 7)
    h1 = out_doc.add_heading("Part 1: Review 1 Rubric & Readiness Checklist", level=1)
    
    # Copy Table 7 from source
    t7_src = src_doc.tables[7]
    t7_new = out_doc.add_table(rows=len(t7_src.rows), cols=len(t7_src.columns))
    t7_new.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, r in enumerate(t7_src.rows):
        for c_i, c in enumerate(r.cells):
            cell = t7_new.cell(r_i, c_i)
            cell.text = c.text
            p = cell.paragraphs[0]
            if r_i == 0:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(9.5)
            else:
                p.runs[0].font.size = Pt(8.5)
    apply_table_formatting(t7_new)
    out_doc.add_paragraph()

    # 2. Slide-by-Slide Presentation Support Plan (Section 36)
    out_doc.add_heading("Part 2: Review 1 Presentation Support (20-Slide Master Plan)", level=1)
    for i in range(390, 513):
        p_src = src_doc.paragraphs[i]
        txt = p_src.text.strip()
        if not txt:
            continue
        sname = p_src.style.name if p_src.style else ''
        if sname.startswith("Heading 1"):
            continue
        elif sname.startswith("Heading 2"):
            out_doc.add_heading(txt, level=2)
        else:
            out_doc.add_paragraph(txt)

    # 3. Comprehensive Q&A Viva Preparation (Section 37)
    out_doc.add_heading("Part 3: Review 1 Viva Examination Q&A Preparation", level=1)
    for i in range(513, 577):
        p_src = src_doc.paragraphs[i]
        txt = p_src.text.strip()
        if not txt:
            continue
        sname = p_src.style.name if p_src.style else ''
        if sname.startswith("Heading 1"):
            continue
        elif sname.startswith("Heading 2"):
            out_doc.add_heading(txt, level=2)
        else:
            p = out_doc.add_paragraph(txt)
            if txt.startswith("Q:"):
                p.runs[0].bold = True
                p.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    # 4. Final Quality-Control Self-Audit & Verification Sign-Off (Sections 38 & 39)
    out_doc.add_heading("Part 4: Quality-Control Self-Audit & Pre-Submission Sign-Off", level=1)
    for i in range(577, 625):
        p_src = src_doc.paragraphs[i]
        txt = p_src.text.strip()
        if not txt:
            continue
        sname = p_src.style.name if p_src.style else ''
        if sname.startswith("Heading 1"):
            out_doc.add_heading(txt, level=2)
        elif sname.startswith("Heading 2"):
            out_doc.add_heading(txt, level=3)
        else:
            out_doc.add_paragraph(txt)

    # 5. Review 2 Implementation Work Sequence (Section 40)
    out_doc.add_heading("Part 5: Review 2 Implementation Work Sequence (Phase Roadmap)", level=1)
    for i in range(625, 634):
        p_src = src_doc.paragraphs[i]
        txt = p_src.text.strip()
        sname = p_src.style.name if p_src.style else ''
        if not txt or sname.startswith("Heading 1"):
            continue
        out_doc.add_paragraph(txt)

    out_file = 'Review_1_Viva_and_Prep_Guide.docx'
    out_doc.save(out_file)
    print(f"Successfully generated {out_file} containing all internal viva preparation materials!")

if __name__ == '__main__':
    create_prep_guide()
